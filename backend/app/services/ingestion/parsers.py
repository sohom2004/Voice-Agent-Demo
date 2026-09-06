from __future__ import annotations

import math
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


@dataclass
class DocumentSection:
    id: str
    title: str | None = None
    type: str = "paragraph"
    content: str = ""
    children: list["DocumentSection"] = field(default_factory=list)
    source: dict[str, Any] = field(default_factory=dict)


@dataclass
class CanonicalDocument:
    document_id: str
    workspace_id: str
    name: str
    file_type: str
    sections: list[DocumentSection]
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass
class DocumentChunk:
    workspace_id: str
    document_id: str
    document_name: str
    section_path: list[str]
    chunk_index: int
    content: str
    metadata: dict[str, Any]


def _parse_text_markdown(buffer: bytes, file_name: str) -> CanonicalDocument:
    text = buffer.decode("utf-8", errors="replace")
    sections: list[DocumentSection] = []
    current_section: DocumentSection | None = None

    for line in text.splitlines():
        trimmed = line.strip()
        if not trimmed:
            continue
        heading_match = re.match(r"^(#{1,6})\s+(.+)$", trimmed)
        if heading_match:
            current_section = DocumentSection(
                id=f"sec_{len(sections)}",
                title=heading_match.group(2),
                type="section",
                content="",
            )
            sections.append(current_section)
        else:
            para = DocumentSection(id=f"para_{len(sections)}", type="paragraph", content=trimmed)
            if current_section is not None:
                current_section.children.append(para)
            else:
                sections.append(para)

    return CanonicalDocument("", "", file_name, "markdown", sections)


def _parse_pdf(buffer: bytes, file_name: str) -> CanonicalDocument:
    from pypdf import PdfReader
    from io import BytesIO

    reader = PdfReader(BytesIO(buffer))
    text_parts = []
    for page in reader.pages:
        text_parts.append(page.extract_text() or "")
    text = "\n".join(text_parts)
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    sections = [
        DocumentSection(id=f"pdf_{idx}", type="paragraph", content=para, source={"page": 1})
        for idx, para in enumerate(paragraphs)
    ]
    return CanonicalDocument("", "", file_name, "pdf", sections, {"numPages": len(reader.pages)})


def _parse_docx(buffer: bytes, file_name: str) -> CanonicalDocument:
    from io import BytesIO
    from docx import Document

    doc = Document(BytesIO(buffer))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    sections = [
        DocumentSection(id=f"docx_{idx}", type="paragraph", content=para)
        for idx, para in enumerate(paragraphs)
    ]
    return CanonicalDocument("", "", file_name, "docx", sections)


def _parse_spreadsheet(buffer: bytes, file_name: str) -> CanonicalDocument:
    from io import BytesIO
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(buffer), read_only=True, data_only=True)
    sections: list[DocumentSection] = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(cell or "") for cell in rows[0]]
        sheet_section = DocumentSection(
            id=f"sheet_{sheet_name}",
            title=f"Sheet: {sheet_name}",
            type="section",
            content=f"Excel Sheet: {sheet_name}. Columns: {', '.join(headers)}",
            source={"sheet": sheet_name},
        )
        for row_idx, row in enumerate(rows[1:], start=1):
            row_text = "\n".join(f"{headers[i]}: {row[i]}" for i in range(min(len(headers), len(row))))
            sheet_section.children.append(
                DocumentSection(
                    id=f"row_{sheet_name}_{row_idx}",
                    type="table",
                    content=row_text,
                    source={"sheet": sheet_name, "row": row_idx},
                )
            )
        sections.append(sheet_section)
    return CanonicalDocument("", "", file_name, Path(file_name).suffix.lstrip("."), sections)


def parse_document(storage_path: str, file_name: str) -> CanonicalDocument:
    buffer = Path(storage_path).read_bytes()
    ext = Path(file_name).suffix.lower().lstrip(".")

    if ext in {"txt", "md", "markdown", "text"}:
        return _parse_text_markdown(buffer, file_name)
    if ext == "pdf":
        return _parse_pdf(buffer, file_name)
    if ext == "docx":
        return _parse_docx(buffer, file_name)
    if ext in {"xlsx", "xls", "csv"}:
        if ext == "csv":
            text = buffer.decode("utf-8", errors="replace")
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            sections = [
                DocumentSection(id=f"csv_{idx}", type="paragraph", content=line)
                for idx, line in enumerate(lines)
            ]
            return CanonicalDocument("", "", file_name, "csv", sections)
        return _parse_spreadsheet(buffer, file_name)

    raise ValueError(f"Unsupported file type: {ext} for file {file_name}")
